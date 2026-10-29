"""주제(프로젝트)별 채널 찾기·수집·분석 파이프라인.

사용법: python3 tools/pipeline.py <프로젝트 폴더> discover|collect|analyze
프로젝트 폴더의 config.json:
  {"name": "야담", "keyword_regex": "...", "queries": [...], "benchmarks": ["https://www.youtube.com/channel/UC..."]}
"""
import csv, glob, json, os, re, statistics, sys, time, urllib.parse
from collections import Counter, defaultdict
sys.path.insert(0, os.path.dirname(__file__))
from yt_scrape import channel_videos, get, initial_data, ytcfg, walk, text, continuation

P = sys.argv[1]
CFG = json.load(open(f"{P}/config.json", encoding="utf-8"))
RAW = f"{P}/raw"
os.makedirs(RAW, exist_ok=True)
KW = re.compile(CFG["keyword_regex"])
LONG = "EgIYAg%3D%3D"


def known_ids():
    out = {}
    for f in glob.glob(f"{RAW}/[0-9]*.json"):
        cid = json.load(open(f, encoding="utf-8"))["info"].get("channel_id")
        if cid:
            out[cid] = f
    return out


def search(q, pages, sp):
    found = []
    url = "https://www.youtube.com/results?search_query=" + urllib.parse.quote(q) + (f"&sp={sp}" if sp else "")
    html = get(url)
    data = initial_data(html)
    key, ver = ytcfg(html)

    def take(d):
        for v in walk(d, "videoRenderer"):
            runs = (v.get("ownerText") or v.get("longBylineText") or {}).get("runs") or [{}]
            cid = runs[0].get("navigationEndpoint", {}).get("browseEndpoint", {}).get("browseId")
            if cid:
                found.append((cid, runs[0].get("text", ""), text(v.get("title"))))
    take(data)
    tok = continuation(data)
    for _ in range(pages - 1):
        if not tok:
            break
        body = json.dumps({"context": {"client": {"clientName": "WEB", "clientVersion": ver, "hl": "ko", "gl": "KR"}}, "continuation": tok}).encode()
        d = json.loads(get("https://www.youtube.com/youtubei/v1/search?prettyPrint=false" + (f"&key={key}" if key else ""), body))
        take(d)
        tok = continuation(d)
        time.sleep(0.3)
    return found


def discover():
    hits = defaultdict(lambda: {"name": "", "titles": {}, "queries": set()})
    prev = f"{RAW}/_discovered.json"
    if os.path.exists(prev):
        for r in json.load(open(prev, encoding="utf-8")):
            h = hits[r["channel_id"]]; h["name"] = r["name"]
            h["titles"].update({t: 1 for t in r["titles"]}); h["queries"].update(r["queries"])
    queries = CFG["queries"] + CFG.get("extra_queries", [])
    done = {q for h in hits.values() for q in h["queries"]}
    for q in queries:
        if q in done:
            continue
        for sp in (LONG, None):
            try:
                for cid, name, title in search(q, 6, sp):
                    h = hits[cid]; h["name"] = name; h["titles"][title] = 1; h["queries"].add(q)
            except Exception as e:
                print("오류", q, e, flush=True)
        print(q, len(hits), flush=True)
    out = [{"channel_id": c, "name": h["name"], "titles": list(h["titles"]), "queries": sorted(h["queries"]),
            "n_kw": sum(bool(KW.search(t)) for t in h["titles"])} for c, h in hits.items()]
    out.sort(key=lambda r: (-r["n_kw"], -len(r["queries"])))
    json.dump(out, open(prev, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("TOTAL", len(out), "후보(주제 제목 2개 이상)", sum(1 for r in out if r["n_kw"] >= 2))


def collect():
    have = known_ids()
    disc = json.load(open(f"{RAW}/_discovered.json", encoding="utf-8"))
    cands = [r for r in disc if r["n_kw"] >= CFG.get("min_hits", 2) and r["channel_id"] not in have]
    nxt = max([int(os.path.basename(f)[:-5]) for f in glob.glob(f"{RAW}/[0-9]*.json")] + [999]) + 1
    for b in CFG.get("benchmarks", []):
        cid = b.rsplit("/", 1)[-1]
        if cid not in have:
            cands.insert(0, {"channel_id": cid, "name": "", "bench": True})
    print("수집할 채널", len(cands), flush=True)
    for k, r in enumerate(cands):
        link = "https://www.youtube.com/channel/" + r["channel_id"]
        rec = {"list_name": r["name"], "list_category": "기준" if r.get("bench") else "검색으로 발견", "input_link": link}
        try:
            info, vids = channel_videos(link, max_pages=CFG.get("max_pages", 33))
            if info:
                rec.update(info); rec["channel_url"] = link; rec["status"] = "ok"; rec["n_collected"] = len(vids)
                i = 0 if r.get("bench") and not os.path.exists(f"{RAW}/000.json") else nxt
                if i == nxt:
                    nxt += 1
                json.dump({"info": rec, "videos": vids}, open(f"{RAW}/{i:03d}.json", "w", encoding="utf-8"), ensure_ascii=False)
        except Exception as e:
            rec["status"] = "오류 " + str(e)[:80]
        print(k, rec.get("name", r["name"]), rec.get("subscribers", ""), rec.get("n_collected", rec.get("status")), flush=True)


def secs(s):
    t = 0
    for x in (s or "").split(":"):
        if x.isdigit():
            t = t * 60 + int(x)
    return t or None


def views(s):
    s = (s or "").replace("조회수", "").replace("회", "").replace(",", "").strip()
    m = re.match(r"([\d.]+)\s*(천|만|억)?", s)
    return float(m.group(1)) * {"천": 1e3, "만": 1e4, "억": 1e8, None: 1}[m.group(2)] if m else None


def days_ago(s):
    m = re.search(r"(\d+)\s*(초|분|시간|일|주|개월|년)", s or "")
    return int(m.group(1)) * {"초": 0, "분": 0, "시간": 0, "일": 1, "주": 7, "개월": 30, "년": 365}[m.group(2)] if m else None


def fmt_len(sec):
    h, m = divmod(int(sec) // 60, 60)
    return f"{h}시간 {m}분" if h else f"{m}분"


def subs_num(s):
    m = re.search(r"([\d.]+)\s*(천|만)?", s or "")
    return float(m.group(1)) * {"천": 1e3, "만": 1e4, None: 1}[m.group(2)] if m else 0


def analyze():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    hi, lo, ex = CFG.get("grade_hi", 0.6), CFG.get("grade_mid", 0.4), CFG.get("exclude_below", 0.15)
    summary, allv, excluded = [], [], []
    for f in sorted(glob.glob(f"{RAW}/[0-9]*.json")):
        d = json.load(open(f, encoding="utf-8"))
        info, vids = d["info"], d["videos"]
        n = len(vids)
        lens = [secs(v["length"]) for v in vids if secs(v["length"])]
        kw = sum(bool(KW.search(v["title"])) for v in vids) / n if n else 0
        lng = sum(1 for x in lens if x >= 3600) / len(lens) if lens else 0
        mm = CFG.get("match_min_minutes", 50) * 60
        match = sum(1 for v in vids if KW.search(v["title"]) and (secs(v["length"]) or 0) >= mm) / n if n else 0
        rv = [views(v["views"]) for v in vids[:10] if views(v["views"]) is not None]
        recent = [v for v in vids if days_ago(v["published"]) is not None and days_ago(v["published"]) <= 30]
        top = max(vids, key=lambda v: views(v["views"]) or 0) if vids else None
        last = days_ago(vids[0]["published"]) if vids else None
        st = "⚪" if last is None else ("🟢" if last <= 7 else ("🟡" if last <= 30 else "🔴"))
        bench = info.get("list_category") == "기준"
        g = "기준 채널" if bench else ("상" if match >= 0.5 else ("중" if match >= 0.25 or kw >= hi else "하"))
        r = {"채널이름": info.get("name") or info.get("list_name"), "채널링크": info.get("channel_url"),
             "구독자": (info.get("subscribers") or "").replace("구독자 ", "").replace("명", ""),
             "영상 수": (info.get("video_count") or "").replace("동영상 ", "").replace("개", "") or str(n),
             "유사도": g, "운영": st, "같은 종류 영상 비율": f"{match:.0%}", f"{CFG['name']} 소재 비율": f"{kw:.0%}", "1시간 이상 영상 비율": f"{lng:.0%}",
             "평균 길이": fmt_len(statistics.mean(lens)) if lens else "", "최근 30일 업로드": len(recent),
             "최근 업로드": vids[0]["published"] if vids else "", "최근 10개 평균 조회수": f"{statistics.mean(rv):,.0f}" if rv else "",
             "조회수 1위 영상": top["title"] if top else "", "조회수 1위 링크": "https://www.youtube.com/watch?v=" + top["id"] if top else "",
             "조회수 1위 조회수": top["views"] if top else "", "최신 영상 제목": vids[0]["title"] if vids else ""}
        if not bench and n == 0:
            r["제외 이유"] = "공개 영상 없음"; excluded.append(r); continue
        if not bench and kw < ex:
            r["제외 이유"] = f"{CFG['name']} 소재 거의 없음 ({kw:.0%})"; excluded.append(r); continue
        summary.append(r)
        for v in vids:
            allv.append({"채널이름": r["채널이름"], "채널링크": r["채널링크"], "제목": v["title"], "영상 링크": "https://www.youtube.com/watch?v=" + v["id"],
                         "조회수": v["views"], "길이": v["length"], "업로드": v["published"]})
    order = {"기준 채널": 0, "상": 1, "중": 2, "하": 3}
    summary.sort(key=lambda r: (order[r["유사도"]], -subs_num(r["구독자"])))
    name = CFG["name"]
    with open(f"{P}/{name}_채널분석.csv", "w", encoding="utf-8-sig", newline="") as fp:
        w = csv.DictWriter(fp, fieldnames=list(summary[0].keys())); w.writeheader(); w.writerows(summary)
    with open(f"{P}/{name}_전체영상제목.csv", "w", encoding="utf-8-sig", newline="") as fp:
        w = csv.DictWriter(fp, fieldnames=list(allv[0].keys())); w.writeheader(); w.writerows(allv)
    with open(f"{P}/{name}_제외한채널.csv", "w", encoding="utf-8-sig", newline="") as fp:
        w = csv.writer(fp); w.writerow(["채널이름", "채널링크", "구독자", "제외 이유"])
        for r in excluded:
            w.writerow([r["채널이름"], r["채널링크"], r["구독자"], r["제외 이유"]])
    # 유사도별 엑셀 (기준 채널은 '상' 시트 맨 위)
    cols = ["채널이름", "운영", "구독자", "영상 수", "같은 종류 영상 비율", f"{name} 소재 비율", "1시간 이상 영상 비율", "평균 길이", "최근 30일 업로드", "최근 업로드", "최근 10개 평균 조회수", "조회수 1위 영상"]
    wb = Workbook(); wb.remove(wb.active)
    colors = {"상": "E2EFDA", "중": "FFF2CC", "하": "FCE4D6"}
    for gname in ("상", "중", "하"):
        part = [r for r in summary if r["유사도"] == gname or (gname == "상" and r["유사도"] == "기준 채널")]
        ws = wb.create_sheet(f"{gname} ({len(part)}개)")
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor=colors[gname])
        for r in part:
            ws.append([r[c] for c in cols])
            row = ws.max_row
            ws.cell(row=row, column=1).hyperlink = r["채널링크"]
            ws.cell(row=row, column=1).font = Font(color="0563C1", underline="single", bold=True)
            if r["조회수 1위 링크"]:
                ws.cell(row=row, column=12).hyperlink = r["조회수 1위 링크"]
                ws.cell(row=row, column=12).font = Font(color="0563C1", underline="single")
            if r["유사도"] == "기준 채널":
                ws.cell(row=row, column=1).value = r["채널이름"] + " ★기준"
        for i, wdt in enumerate([30, 6, 10, 9, 16, 13, 16, 12, 14, 12, 16, 70], 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = wdt
        ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
        with open(f"{P}/{name}_유사도_{gname}.csv", "w", encoding="utf-8-sig", newline="") as fp:
            w = csv.writer(fp); w.writerow(["채널이름", "채널링크"] + cols[1:])
            for r in part:
                w.writerow([r["채널이름"], r["채널링크"]] + [r[c] for c in cols[1:]])
    wb.save(f"{P}/{name}_유사도별_채널목록.xlsx")
    print("남은 채널", len(summary), Counter(r["유사도"] for r in summary), "제외", len(excluded), "영상", len(allv))


def related(vid):
    """영상 옆 '관련 영상' 목록: [(영상ID, 채널이름, 제목)]"""
    body = json.dumps({"context": {"client": {"clientName": "WEB", "clientVersion": "2.20241001.00.00", "hl": "ko", "gl": "KR"}}, "videoId": vid}).encode()
    d = json.loads(get("https://www.youtube.com/youtubei/v1/next?prettyPrint=false", body))
    out = []
    for v in walk(d, "lockupViewModel"):
        if v.get("contentType") != "LOCKUP_CONTENT_TYPE_VIDEO":
            continue
        meta = v.get("metadata", {}).get("lockupMetadataViewModel", {})
        rows = [text(p.get("text")) for p in walk(meta, "metadataParts") for p in (p if isinstance(p, list) else [p])]
        if rows:
            out.append((v.get("contentId"), rows[0], text(meta.get("title"))))
    return out


def channel_of(vid):
    """영상 ID로 채널 주소 찾기 (oEmbed)"""
    try:
        d = json.loads(get("https://www.youtube.com/oembed?format=json&url=" + urllib.parse.quote("https://www.youtube.com/watch?v=" + vid)))
        return d.get("author_url")
    except Exception:
        return None


def crawl():
    """기준·상위 채널 영상의 관련 영상을 따라가며 비슷한 채널을 찾는다."""
    from concurrent.futures import ThreadPoolExecutor
    state_f = f"{RAW}/_crawl.json"
    st = json.load(open(state_f, encoding="utf-8")) if os.path.exists(state_f) else {"seen_videos": [], "channels": {}}
    seen = set(st["seen_videos"])
    chans = st["channels"]  # 채널이름 -> {"titles": {영상ID: 제목}, "url": None}
    have_names = set()
    seeds = []
    for f in glob.glob(f"{RAW}/[0-9]*.json"):
        d = json.load(open(f, encoding="utf-8"))
        vids = d["videos"]
        mm = CFG.get("match_min_minutes", 50) * 60
        match = sum(1 for v in vids if KW.search(v["title"]) and (secs(v["length"]) or 0) >= mm) / len(vids) if vids else 0
        have_names.add(d["info"].get("name"))
        if d["info"].get("list_category") == "기준" or match >= 0.5:
            top = sorted(vids, key=lambda v: -(views(v["views"]) or 0))[:CFG.get("crawl_per_channel", 8)]
            seeds += [v["id"] for v in top + vids[:3]]
    seeds = [v for v in dict.fromkeys(seeds) if v not in seen][:CFG.get("crawl_max", 400)]
    print("관련 영상 조사할 영상", len(seeds), flush=True)

    def one(vid):
        try:
            return vid, related(vid)
        except Exception:
            return vid, []
    with ThreadPoolExecutor(4) as ex:
        for k, (vid, rel) in enumerate(ex.map(one, seeds), 1):
            seen.add(vid)
            for rid, name, title in rel:
                c = chans.setdefault(name, {"titles": {}, "url": None})
                c["titles"][rid] = title
            if k % 50 == 0:
                print(k, "채널 후보", len(chans), flush=True)
    # 주제 제목이 2개 이상인 새 채널만 주소 찾기
    new = [(n, c) for n, c in chans.items() if n not in have_names and not c["url"]
           and sum(bool(KW.search(t)) for t in c["titles"].values()) >= CFG.get("min_hits", 2)]
    print("주소 찾을 새 채널", len(new), flush=True)
    for n, c in new:
        c["url"] = channel_of(next(iter(c["titles"])))
    st = {"seen_videos": sorted(seen), "channels": chans}
    json.dump(st, open(state_f, "w", encoding="utf-8"), ensure_ascii=False)
    # collect 단계가 읽는 후보 목록(_discovered.json)에 합치기
    disc_f = f"{RAW}/_discovered.json"
    disc = json.load(open(disc_f, encoding="utf-8")) if os.path.exists(disc_f) else []
    by = {r["channel_id"]: r for r in disc}
    for n, c in chans.items():
        if not c.get("url"):
            continue
        cid = c.get("cid")
        if not cid:
            try:
                info, _ = channel_videos(c["url"], max_pages=0)
                cid = c["cid"] = info.get("channel_id") if info else None
            except Exception:
                cid = None
        if cid and cid not in by:
            titles = list(c["titles"].values())
            by[cid] = {"channel_id": cid, "name": n, "titles": titles, "queries": ["관련 영상"], "n_kw": sum(bool(KW.search(t)) for t in titles)}
    json.dump(list(by.values()), open(disc_f, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump({"seen_videos": sorted(seen), "channels": chans}, open(state_f, "w", encoding="utf-8"), ensure_ascii=False)
    print("후보 채널 총", len(by), flush=True)


{"discover": discover, "collect": collect, "analyze": analyze, "crawl": crawl}[sys.argv[2]]()
