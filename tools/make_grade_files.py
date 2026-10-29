"""최종 채널 분석표를 유사도(상/중/하)별로 나눠 보기 쉬운 엑셀과 CSV로 만든다."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

rows = list(csv.DictReader(open("channels/최종_채널분석.csv", encoding="utf-8-sig")))
COLS = ["채널이름", "구독자", "영상 수", "시댁 소재 비율", "1시간 이상 영상 비율", "평균 길이",
        "최근 30일 업로드", "최근 10개 평균 조회수", "조회수 1위 영상", "채널링크", "조회수 1위 링크"]
WIDTH = {"채널이름": 28, "구독자": 10, "영상 수": 9, "시댁 소재 비율": 13, "1시간 이상 영상 비율": 17, "평균 길이": 13,
         "최근 30일 업로드": 15, "최근 10개 평균 조회수": 18, "조회수 1위 영상": 70, "채널링크": 14, "조회수 1위 링크": 14}
GROUPS = [("기준 채널", "기준 채널", "D9E1F2"), ("상", "상 (거의 같음)", "E2EFDA"), ("중", "중 (비슷함)", "FFF2CC"), ("하", "하 (일부만 비슷)", "FCE4D6")]

wb = Workbook()
guide = wb.active
guide.title = "읽는 법"
lines = [
    ("세상만사·사연만남과 비슷한 채널 목록", True),
    ("", False),
    ("시트 구성", True),
    (f"· 기준 채널: 세상만사, 사연만남 ({sum(r['유사도']=='기준 채널' for r in rows)}개)", False),
    (f"· 상 (거의 같음): 시댁 소재 60% 이상 + 1시간 넘는 영상 50% 이상 ({sum(r['유사도']=='상' for r in rows)}개)", False),
    (f"· 중 (비슷함): 시댁 소재 40% 이상이지만 영상이 짧거나 다른 사연이 섞임 ({sum(r['유사도']=='중' for r in rows)}개)", False),
    (f"· 하 (일부만 비슷): 시댁 소재 15~40% ({sum(r['유사도']=='하' for r in rows)}개)", False),
    ("", False),
    ("열 설명", True),
    ("· 채널이름을 누르면 YouTube 채널로 이동합니다.", False),
    ("· 시댁 소재 비율: 영상 제목에 시모·시누·동서·며느리 등이 들어간 비율", False),
    ("· 1시간 이상 영상 비율: 전체 영상 중 1시간 넘는 영상의 비율", False),
    ("· 최근 30일 업로드: 최근 한 달 동안 올린 영상 수 (0이면 업로드를 멈춘 채널)", False),
    ("· 조회수 1위 영상을 누르면 그 영상으로 이동합니다.", False),
    ("· 각 시트는 구독자 많은 순서입니다.", False),
    ("", False),
    ("기준일: 2026-10-29 (쇼츠·라이브 제외)", False),
]
for t, bold in lines:
    guide.append([t])
    guide.cell(row=guide.max_row, column=1).font = Font(bold=bold, size=14 if bold else 11)
guide.column_dimensions["A"].width = 100

for key, title, color in GROUPS:
    part = [r for r in rows if r["유사도"] == key]
    ws = wb.create_sheet(title.split(" ")[0] if key != "기준 채널" else "기준 채널")
    ws.append(COLS[:-2])
    fill = PatternFill("solid", fgColor=color)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in part:
        ws.append([r[c] for c in COLS[:-2]])
        row = ws.max_row
        name = ws.cell(row=row, column=1)
        name.hyperlink = r["채널링크"]
        name.font = Font(color="0563C1", underline="single", bold=True)
        top = ws.cell(row=row, column=9)
        if r["조회수 1위 링크"]:
            top.hyperlink = r["조회수 1위 링크"]
            top.font = Font(color="0563C1", underline="single")
    for i, c in enumerate(COLS[:-2], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = WIDTH[c]
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    with open(f"channels/유사도_{key.replace(' 채널', '')}.csv", "w", encoding="utf-8-sig", newline="") as fp:
        w = csv.writer(fp)
        w.writerow(["채널이름", "채널링크"] + COLS[1:9] + ["조회수 1위 링크"])
        for r in part:
            w.writerow([r["채널이름"], r["채널링크"]] + [r[c] for c in COLS[1:9]] + [r["조회수 1위 링크"]])
wb.save("channels/유사도별_채널목록.xlsx")
print("ok")
