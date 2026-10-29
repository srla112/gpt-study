"""상 채널 영상 제목 중 완전히 같은 제목과 거의 같은 제목(같은 사연을 조금 바꾼 것)을 찾는다."""
import csv, re
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

rows = list(csv.DictReader(open("channels/상_채널_영상제목_조회수.csv", encoding="utf-8-sig")))
link = {}
for v in csv.DictReader(open("channels/전체_영상제목.csv", encoding="utf-8-sig")):
    link[(v["채널이름"].split(" (@")[0], v["제목"])] = v["영상 링크"]
for i, r in enumerate(rows):
    r["i"] = i
    r["제목"] = r["영상 제목"]
    r["조회수"] = int(r["조회수"]) if r["조회수"] else 0
    r["링크"] = link.get((r["채널이름"].split(" (@")[0], r["제목"]), "")


def norm(t):
    t = re.sub(r"#\S+", "", t)
    t = re.sub(r"\[[^\]]*\]|【[^】]*】", "", t)
    t = re.sub(r"[^\w가-힣]", "", t)
    return t.lower()


for r in rows:
    r["n"] = norm(r["제목"])

# 1) 완전히 같은 제목 (태그·띄어쓰기·문장부호 무시)
exact = defaultdict(list)
for r in rows:
    if len(r["n"]) >= 10:
        exact[r["n"]].append(r)
exact_groups = [g for g in exact.values() if len(g) > 1]

# 2) 거의 같은 제목: 글자 4개 묶음이 65% 이상 겹치면 같은 사연으로 본다
grams = [set(r["n"][k:k + 4] for k in range(len(r["n"]) - 3)) for r in rows]
df = Counter(g for s in grams for g in s)
inv = defaultdict(list)
for i, s in enumerate(grams):
    for g in s:
        if df[g] <= 60:
            inv[g].append(i)
parent = list(range(len(rows)))


def find(x):
    while parent[x] != x:
        parent[x] = parent[parent[x]]
        x = parent[x]
    return x


for i, s in enumerate(grams):
    if len(s) < 8:
        continue
    cnt = Counter(j for g in s if df[g] <= 60 for j in inv[g] if j > i)
    for j, c in cnt.items():
        if c >= 5:
            jac = len(s & grams[j]) / len(s | grams[j])
            if jac >= 0.65:
                parent[find(i)] = find(j)
clusters = defaultdict(list)
for i in range(len(rows)):
    clusters[find(i)].append(rows[i])
near_groups = [g for g in clusters.values() if len(g) > 1]
near_groups.sort(key=lambda g: (-len({r["채널이름"] for r in g}), -len(g)))

# 3) 채널끼리 같은 사연을 얼마나 공유하는지
pair = Counter()
for g in near_groups:
    chs = sorted({r["채널이름"] for r in g})
    for a in range(len(chs)):
        for b in range(a + 1, len(chs)):
            pair[(chs[a], chs[b])] += 1
self_dup = Counter()
for g in near_groups:
    c = Counter(r["채널이름"] for r in g)
    for ch, k in c.items():
        if k > 1:
            self_dup[ch] += k - 1

H = Font(bold=True)
F = PatternFill("solid", fgColor="E2EFDA")
A = PatternFill("solid", fgColor="F2F2F2")


def header(ws, cols, widths):
    ws.append(cols)
    for c in ws[1]:
        c.font = H; c.fill = F
    for k, w in enumerate(widths):
        ws.column_dimensions[chr(65 + k)].width = w
    ws.freeze_panes = "A2"


wb = Workbook()
s = wb.active
s.title = "요약"
multi = [g for g in near_groups if len({r["채널이름"] for r in g}) > 1]
for line in [
    "상 채널 83개 영상 제목 중복 분석 (영상 18,683개)", "",
    f"· 완전히 같은 제목 묶음: {len(exact_groups)}개 (영상 {sum(len(g) for g in exact_groups)}개)",
    f"· 거의 같은 제목(같은 사연) 묶음: {len(near_groups)}개 (영상 {sum(len(g) for g in near_groups)}개)",
    f"   - 그중 서로 다른 채널끼리 겹치는 묶음: {len(multi)}개",
    f"   - 같은 채널 안에서만 겹치는 묶음(재업로드): {len(near_groups) - len(multi)}개", "",
    "판단 기준", "· 완전히 같음: 태그([사이다사연] 등)·해시태그·띄어쓰기·문장부호를 빼고 글자가 똑같은 제목",
    "· 거의 같음: 제목 글자의 65% 이상이 겹치는 제목 (단어 몇 개만 바꾼 같은 사연)",
]:
    s.append([line])
s["A1"].font = Font(bold=True, size=14)
s.column_dimensions["A"].width = 100

ws = wb.create_sheet("완전히 같은 제목")
header(ws, ["묶음", "채널이름", "영상 제목", "조회수", "영상 링크"], [7, 28, 100, 11, 14])
for n, g in enumerate(sorted(exact_groups, key=lambda g: -len(g)), 1):
    for r in sorted(g, key=lambda r: -r["조회수"]):
        ws.append([n, r["채널이름"], r["제목"], r["조회수"], r["링크"]])
        if n % 2 == 0:
            for c in ws[ws.max_row]:
                c.fill = A
        if r["링크"]:
            ws.cell(row=ws.max_row, column=5).hyperlink = r["링크"]
        ws.cell(row=ws.max_row, column=4).number_format = "#,##0"

ws = wb.create_sheet("거의 같은 제목")
header(ws, ["묶음", "채널 수", "채널이름", "영상 제목", "조회수", "영상 링크"], [7, 8, 28, 100, 11, 14])
for n, g in enumerate(near_groups, 1):
    nch = len({r["채널이름"] for r in g})
    for r in sorted(g, key=lambda r: -r["조회수"]):
        ws.append([n, nch, r["채널이름"], r["제목"], r["조회수"], r["링크"]])
        if n % 2 == 0:
            for c in ws[ws.max_row]:
                c.fill = A
        if r["링크"]:
            ws.cell(row=ws.max_row, column=6).hyperlink = r["링크"]
        ws.cell(row=ws.max_row, column=5).number_format = "#,##0"
ws.auto_filter.ref = ws.dimensions

ws = wb.create_sheet("채널끼리 겹침")
header(ws, ["채널 A", "채널 B", "같은 사연 수"], [32, 32, 14])
for (a, b), k in pair.most_common():
    ws.append([a, b, k])

ws = wb.create_sheet("채널 안 재업로드")
header(ws, ["채널이름", "같은 사연을 다시 올린 횟수"], [32, 24])
for ch, k in self_dup.most_common():
    ws.append([ch, k])

wb.save("channels/상_채널_중복제목.xlsx")
print("exact", len(exact_groups), sum(len(g) for g in exact_groups))
print("near", len(near_groups), sum(len(g) for g in near_groups), "multi", len(multi))
print("top pairs", pair.most_common(12))
print("self", self_dup.most_common(10))
for g in multi[:6]:
    print("---")
    for r in g[:4]:
        print(r["채널이름"], "|", r["제목"][:90], "|", r["조회수"])

for base in ("세상만사", "사연만남 (@사연만남1짱)"):
    gs = [g for g in near_groups if any(r["채널이름"] == base for r in g)]
    other = [g for g in gs if any(r["채널이름"] != base for r in g)]
    tot = sum(1 for r in rows if r["채널이름"] == base)
    print(base, "전체", tot, "겹치는 묶음", len(gs), "다른 채널과 겹침", len(other))
    for g in other[:5]:
        print("  ", [(r["채널이름"], r["조회수"]) for r in g][:5], g[0]["제목"][:60])
