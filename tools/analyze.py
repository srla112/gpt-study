"""수집한 채널/영상 데이터를 분석해 노션용 CSV를 만든다."""
import csv, glob, json, re, statistics

SIDAE = re.compile(r"시모|시어머니|시어머님|시누|시동생|동서|시댁|시부|시아버지|시아버님|시가|시월드|며느리|고부|형님|올케|아랫동서|예비시")
QUOTE = re.compile(r"[\"“”'‘’]")


def secs(s):
    p = [int(x) for x in s.split(":") if x.isdigit()] if s else []
    if not p:
        return None
    t = 0
    for x in p:
        t = t * 60 + x
    return t


def views(s):
    s = (s or "").replace("조회수", "").replace("회", "").replace(",", "").strip()
    m = re.match(r"([\d.]+)\s*(천|만|억)?", s)
    if not m:
        return None
    return float(m.group(1)) * {"천": 1e3, "만": 1e4, "억": 1e8, None: 1}[m.group(2)]


def days_ago(s):
    m = re.search(r"(\d+)\s*(초|분|시간|일|주|개월|년)", s or "")
    if not m:
        return None
    return int(m.group(1)) * {"초": 0, "분": 0, "시간": 0, "일": 1, "주": 7, "개월": 30, "년": 365}[m.group(2)]


def fmt_len(sec):
    if sec is None:
        return ""
    h, m = divmod(int(sec) // 60, 60)
    return f"{h}시간 {m}분" if h else f"{m}분"


def grade(sidae, long_, quote):
    score = sidae * 0.5 + long_ * 0.3 + quote * 0.2
    if sidae >= 0.6 and long_ >= 0.5:
        return "상", score
    if sidae >= 0.4:
        return "중", score
    return "하", score


idx = json.load(open("channels/raw/_index.json", encoding="utf-8"))
summary, allv = [], []
for f in sorted(glob.glob("channels/raw/[0-9]*.json")):
    d = json.load(open(f, encoding="utf-8"))
    info, vids = d["info"], d["videos"]
    n = len(vids)
    lens = [secs(v["length"]) for v in vids if secs(v["length"])]
    recent = [v for v in vids if (days_ago(v["published"]) is not None and days_ago(v["published"]) <= 30)]
    sid = sum(bool(SIDAE.search(v["title"])) for v in vids) / n if n else 0
    lng = sum(1 for x in lens if x >= 3600) / len(lens) if lens else 0
    quo = sum(bool(QUOTE.search(v["title"])) for v in vids) / n if n else 0
    rv = [views(v["views"]) for v in vids[:10] if views(v["views"]) is not None]
    g, sc = grade(sid, lng, quo) if n else ("-", 0)
    top = max(vids, key=lambda v: views(v["views"]) or 0) if vids else None
    summary.append({
        "채널이름": (info.get("name") + " (@사연만남1짱)") if "사연만남1짱" in info["list_name"] else (info.get("name") or info["list_name"]),
        "채널링크": info.get("channel_url", info["input_link"]),
        "구독자": (info.get("subscribers") or "").replace("구독자 ", "").replace("명", ""),
        "영상 수": (info.get("video_count") or "").replace("동영상 ", "").replace("개", "") or str(n),
        "유사도": g,
        "시댁 소재 비율": f"{sid:.0%}",
        "1시간 이상 영상 비율": f"{lng:.0%}",
        "평균 길이": fmt_len(statistics.mean(lens)) if lens else "",
        "최근 30일 업로드": len(recent),
        "최근 업로드": vids[0]["published"] if vids else "",
        "최근 10개 평균 조회수": f"{statistics.mean(rv):,.0f}" if rv else "",
        "조회수 1위 영상": top["title"] if top else "",
        "조회수 1위 링크": "https://www.youtube.com/watch?v=" + top["id"] if top else "",
        "조회수 1위 조회수": top["views"] if top else "",
        "최신 영상 제목": vids[0]["title"] if vids else "",
        "기존 분류": info["list_category"],
        "_score": sc, "_n": n, "_sid": sid,
    })
    for v in vids:
        allv.append({"채널이름": summary[-1]["채널이름"], "제목": v["title"], "영상 링크": "https://www.youtube.com/watch?v=" + v["id"],
                     "조회수": v["views"], "길이": v["length"], "업로드": v["published"],
                     "시댁 소재": "O" if SIDAE.search(v["title"]) else ""})
fails = [r for r in idx if r.get("status") != "ok"]
excluded = []
keep = []
for r in summary:
    base = r["기존 분류"] == "기준"
    if not base and r["_n"] == 0:
        r["제외 이유"] = "공개 영상 없음"; excluded.append(r)
    elif not base and r["_sid"] < 0.15:
        r["제외 이유"] = f"시댁 소재 거의 없음 ({r['시댁 소재 비율']})"; excluded.append(r)
    else:
        keep.append(r)
summary = keep
order = {"기준": 0, "상": 1, "중": 2, "하": 3, "-": 4}
summary.sort(key=lambda r: (0 if r["기존 분류"] == "기준" else order[r["유사도"]], -r["_score"]))
for r in summary:
    if r["기존 분류"] == "기준":
        r["유사도"] = "기준 채널"
    for k in ("_score", "_n", "_sid"):
        del r[k]
with open("channels/최종_채널분석.csv", "w", encoding="utf-8-sig", newline="") as fp:
    w = csv.DictWriter(fp, fieldnames=list(summary[0].keys())); w.writeheader(); w.writerows(summary)
with open("channels/전체_영상제목.csv", "w", encoding="utf-8-sig", newline="") as fp:
    w = csv.DictWriter(fp, fieldnames=list(allv[0].keys())); w.writeheader(); w.writerows(allv)
for r in excluded:
    for k in ("_score", "_n", "_sid"):
        r.pop(k, None)
with open("channels/제외한_채널.csv", "w", encoding="utf-8-sig", newline="") as fp:
    w = csv.writer(fp); w.writerow(["채널이름", "채널링크", "구독자", "제외 이유", "기존 분류"])
    for r in excluded:
        w.writerow([r["채널이름"], r["채널링크"], r["구독자"], r["제외 이유"], r["기존 분류"]])
    for r in fails:
        w.writerow([r["list_name"], r["input_link"], "", "채널을 찾지 못함: " + str(r.get("status")), r["list_category"]])
with open("channels/찾지못한_채널.csv", "w", encoding="utf-8-sig", newline="") as fp:
    w = csv.writer(fp); w.writerow(["목록 이름", "분류", "원래 링크", "상태", "검색 후보"])
    for r in fails:
        w.writerow([r["list_name"], r["list_category"], r["input_link"], r.get("status"), ", ".join(r.get("search_candidates", []))])
print(len(summary), len(allv), len(excluded), len(fails))

from openpyxl import Workbook
from openpyxl.styles import Font
wb = Workbook()
for title, path in [("채널 분석", "channels/최종_채널분석.csv"), ("전체 영상 제목", "channels/전체_영상제목.csv"), ("제외한 채널", "channels/제외한_채널.csv")]:
    ws = wb.create_sheet(title)
    rows = list(csv.reader(open(path, encoding="utf-8-sig")))
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
    hdr = rows[0]
    for col, name in enumerate(hdr, 1):
        if "링크" in name:
            for row in range(2, len(rows) + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value and str(cell.value).startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
del wb["Sheet"]
wb.save("channels/최종_채널분석.xlsx")
