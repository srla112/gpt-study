"""유사도 '상' 채널(기준 채널 포함)의 영상 제목과 조회수를 채널별 시트로 정리한다."""
import csv, re
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

top = list(csv.DictReader(open("channels/유사도_상.csv", encoding="utf-8-sig")))
link2name = {r["채널링크"]: r["채널이름"] for r in top}
an = {r["채널이름"]: r for r in csv.DictReader(open("channels/최종_채널분석.csv", encoding="utf-8-sig"))}
vids = defaultdict(list)
for v in csv.DictReader(open("channels/전체_영상제목.csv", encoding="utf-8-sig")):
    vids[v["채널이름"]].append(v)


def views(s):
    s = (s or "").replace("조회수", "").replace("회", "").replace(",", "").strip()
    m = re.match(r"([\d.]+)\s*(천|만|억)?", s)
    return int(round(float(m.group(1)) * {"천": 1e3, "만": 1e4, "억": 1e8, None: 1}[m.group(2)])) if m else None


wb = Workbook()
idx = wb.active
idx.title = "목차"
idx.append(["번호", "채널이름 (누르면 해당 시트로 이동)", "구독자", "정리한 영상 수"])
for c in idx[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="E2EFDA")
used = set()
allrows = []
for n, r in enumerate(top, 1):
    name = r["채널이름"].replace(" ★기준", "")
    key = an.get(name, {}).get("채널이름", name)
    items = vids.get(key, [])
    sheet = re.sub(r"[\[\]:*?/\\']", "", name)[:24].strip() or f"채널{n}"
    base = sheet
    k = 2
    while sheet in used:
        sheet = f"{base[:21]}_{k}"; k += 1
    used.add(sheet)
    ws = wb.create_sheet(sheet)
    ws.append([f"{name}  (구독자 {r['구독자']})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws["A1"].hyperlink = r["채널링크"]
    ws.append(["번호", "영상 제목", "조회수"])
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E2EFDA")
    for i, v in enumerate(items, 1):
        ws.append([i, v["제목"], views(v["조회수"])])
        ws.cell(row=ws.max_row, column=3).number_format = "#,##0"
        allrows.append([name, v["제목"], views(v["조회수"])])
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 110
    ws.column_dimensions["C"].width = 12
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:C{ws.max_row}"
    idx.append([n, name, r["구독자"], len(items)])
    cell = idx.cell(row=idx.max_row, column=2)
    cell.hyperlink = f"#'{sheet}'!A1"
    cell.font = Font(color="0563C1", underline="single")
idx.column_dimensions["B"].width = 45
idx.column_dimensions["C"].width = 12
idx.column_dimensions["D"].width = 15
idx.freeze_panes = "A2"
wb.move_sheet("목차", offset=-len(wb.sheetnames))
allws = wb.create_sheet("전체 모아보기", 1)
allws.append(["채널이름", "영상 제목", "조회수"])
for c in allws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="E2EFDA")
for row in allrows:
    allws.append(row)
    allws.cell(row=allws.max_row, column=3).number_format = "#,##0"
allws.column_dimensions["A"].width = 28
allws.column_dimensions["B"].width = 110
allws.column_dimensions["C"].width = 12
allws.freeze_panes = "A2"
allws.auto_filter.ref = allws.dimensions
wb.save("channels/상_채널_영상제목_조회수.xlsx")
with open("channels/상_채널_영상제목_조회수.csv", "w", encoding="utf-8-sig", newline="") as fp:
    w = csv.writer(fp); w.writerow(["채널이름", "영상 제목", "조회수"]); w.writerows(allrows)
print(len(top), len(allrows), sum(1 for r in top if not vids.get(an.get(r["채널이름"].replace(" ★기준", ""), {}).get("채널이름", r["채널이름"].replace(" ★기준", "")))))
